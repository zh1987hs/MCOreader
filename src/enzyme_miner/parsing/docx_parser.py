import pathlib
from typing import Tuple

from docx import Document


def parse_docx(path: pathlib.Path) -> Tuple[str, dict]:
    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    return text, {"parser": "python-docx", "paragraphs": len(doc.paragraphs)}
